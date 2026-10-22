#!/usr/bin/env python3
"""
Build an editable PowerPoint deck from deck.js.

    python3 -m venv .venv && .venv/bin/pip install python-pptx
    .venv/bin/python build-pptx.py

Writes vLLM-Semantic-Router.pptx next to this script. Every slide is real
text in real text frames — no images — so the result can be restyled in
PowerPoint or Google Slides, or fed back through a deck generator.

deck.js is a JS file only so that index.html can load it with a plain <script>
tag and still work from file://. The object literal inside it is strict JSON,
which is what lets this script read the same source.
"""

import json
import pathlib
import re
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("python-pptx is not installed. See the docstring at the top of this file.")

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "vLLM-Semantic-Router.pptx"

# ----------------------------------------------------------------- palette

BG = RGBColor(0x0B, 0x0E, 0x14)
PANEL = RGBColor(0x15, 0x1B, 0x26)
PANEL2 = RGBColor(0x1E, 0x26, 0x36)
BORDER = RGBColor(0x2E, 0x38, 0x4A)
TEXT = RGBColor(0xEE, 0xF3, 0xFA)
DIM = RGBColor(0xA3, 0xB0, 0xC6)
FAINT = RGBColor(0x6E, 0x7D, 0x96)
ACCENT = RGBColor(0xFF, 0x4D, 0x4D)
ACCENT_DEEP = RGBColor(0xEE, 0x00, 0x00)
CYAN = RGBColor(0x4D, 0xD6, 0xC8)
VIOLET = RGBColor(0xB3, 0x88, 0xFF)
GOLD = RGBColor(0xFF, 0xC8, 0x57)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Helvetica Neue"
MONO = "Menlo"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.72)                      # side margin
CONTENT_W = W - 2 * M
ACCENT_COLS = [ACCENT, CYAN, VIOLET]

# ------------------------------------------------------------- deck.js I/O


def load_deck():
    """Pull the JSON object literal out of deck.js."""
    src = (HERE / "deck.js").read_text(encoding="utf-8")
    m = re.search(r"const\s+DECK\s*=\s*(\{.*?\n\});", src, re.S)
    if not m:
        sys.exit("Could not find `const DECK = {...};` in deck.js")
    try:
        return json.loads(m.group(1))
    except json.JSONDecodeError as e:
        sys.exit(
            f"deck.js is not strict JSON: {e}\n"
            "Keep it double-quoted with no trailing commas — see the header comment."
        )


# ----------------------------------------------------------------- helpers


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return s


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE

    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, color=TEXT, bold=False, font=SANS,
         space_after=0, space_before=0, line=None, first=False, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    if align:
        p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return p


def heading(slide, title, intro=None):
    """Title, optional intro, plus the red top rule. Returns the next free y."""
    rect(slide, Inches(0), Inches(0), Inches(1.95), Pt(5), fill=ACCENT_DEEP)

    tf = textbox(slide, M, Inches(0.58), CONTENT_W, Inches(0.95))
    para(tf, title, 29, TEXT, bold=True, first=True, line=1.1)
    y = Inches(0.58) + Inches(0.62) * (1 + title.count("\n"))

    if intro:
        itf = textbox(slide, M, y + Inches(0.1), Inches(10.6), Inches(0.9))
        para(itf, intro, 14, DIM, first=True, line=1.35)
        y += Inches(0.14) + Inches(0.26) * max(1, round(len(intro) / 112) + 1)

    return y + Inches(0.26)


def chrome(slide, source, idx, total, at=None):
    """Source note bottom-left, clock position + page number bottom-right."""
    if source:
        tf = textbox(slide, M, H - Inches(0.52), Inches(7), Inches(0.3))
        para(tf, source, 9, FAINT, font=MONO, first=True)
    tf = textbox(slide, W - M - Inches(2.2), H - Inches(0.52), Inches(2.2), Inches(0.3))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    if at:
        r = p.add_run()
        r.text = f"{at}  ·  "
        r.font.size, r.font.color.rgb, r.font.name = Pt(10), ACCENT, MONO
    r = p.add_run()
    r.text = f"{idx} / {total}"
    r.font.size, r.font.color.rgb, r.font.name = Pt(10), FAINT, MONO


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def callout(slide, y, text):
    h = Inches(0.52) + Inches(0.2) * (len(text) // 110)
    bar = rect(slide, M, y, CONTENT_W, h, fill=PANEL2, line=GOLD, line_w=1.0)
    rect(slide, M, y, Pt(4), h, fill=GOLD)
    tf = textbox(slide, M + Inches(0.22), y + Inches(0.13), CONTENT_W - Inches(0.44),
                 h - Inches(0.26))
    para(tf, text, 12.5, GOLD, first=True, line=1.3)
    return bar


# ------------------------------------------------------------ slide types


def s_title(slide, sl):
    tf = textbox(slide, M, Inches(1.75), Inches(10.4), Inches(0.4))
    para(tf, sl["subtitle"].upper(), 12, ACCENT, bold=True, font=MONO, first=True)

    tf = textbox(slide, M, Inches(2.3), Inches(10.6), Inches(1.9))
    para(tf, sl["title"], 46, TEXT, bold=True, first=True, line=1.03)

    rect(slide, M, Inches(4.45), Inches(1.1), Pt(6), fill=ACCENT_DEEP)

    tf = textbox(slide, M, Inches(4.85), Inches(8.3), Inches(1.3))
    para(tf, sl["tagline"], 15, DIM, first=True, line=1.45)


def s_closer(slide, sl):
    box_w, box_h = Inches(7.6), Inches(1.6)
    x = (W - box_w) / 2
    rect(slide, x, Inches(2.1), box_w, box_h, fill=PANEL2, line=BORDER, line_w=1.5)
    tf = textbox(slide, x, Inches(2.1), box_w, box_h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, f"> {sl['title']}", 40, TEXT, bold=True, font=MONO, first=True,
         align=PP_ALIGN.CENTER)

    tf = textbox(slide, M, Inches(4.0), CONTENT_W, Inches(0.5))
    para(tf, sl["sub"], 16, DIM, bold=True, first=True, align=PP_ALIGN.CENTER)

    tf = textbox(slide, M, Inches(4.75), CONTENT_W, Inches(1.4))
    for i, link in enumerate(sl.get("links", [])):
        para(tf, link, 12, FAINT, font=MONO, first=(i == 0), space_after=7,
             align=PP_ALIGN.CENTER)


def s_statement(slide, sl, y):
    lead = sl["lead"]
    lh = Inches(0.46) + Inches(0.3) * (len(lead) // 88)
    rect(slide, M, y, Pt(4), lh, fill=ACCENT)
    tf = textbox(slide, M + Inches(0.26), y, Inches(10.9), lh)
    para(tf, lead, 18, TEXT, bold=True, first=True, line=1.3)
    y += lh + Inches(0.34)

    tf = textbox(slide, M, y, Inches(11.2), Inches(3))
    for i, pt in enumerate(sl.get("points", [])):
        p = para(tf, pt, 14.5, DIM, first=(i == 0), space_after=11, line=1.35)
        p.level = 0
        # Manual bullet: a small square glyph keeps the look consistent with
        # the HTML deck without fighting PowerPoint's list styling.
        p.text = "▪   " + pt
        for r in p.runs:
            r.font.size = Pt(14.5)
            r.font.color.rgb = DIM
            r.font.name = SANS
        p.runs[0].font.color.rgb = ACCENT


def s_columns(slide, sl, y):
    cols = sl.get("columns", [])
    n = len(cols) or 1
    gap = Inches(0.22)
    cw = (CONTENT_W - gap * (n - 1)) / n
    avail = H - y - Inches(0.95)
    ch = min(Inches(3.1), avail) if not sl.get("callout") else min(Inches(2.45), avail)

    for i, c in enumerate(cols):
        x = M + (cw + gap) * i
        rect(slide, x, y, cw, ch, fill=PANEL, line=BORDER)
        rect(slide, x, y, cw, Pt(3), fill=ACCENT_COLS[i % len(ACCENT_COLS)])

        tf = textbox(slide, x + Inches(0.22), y + Inches(0.26),
                     cw - Inches(0.44), Inches(0.5))
        para(tf, c["heading"], 14.5, TEXT, bold=True, first=True, line=1.15)

        tf = textbox(slide, x + Inches(0.22), y + Inches(0.78),
                     cw - Inches(0.44), ch - Inches(1.0))
        para(tf, c["body"], 11.5, DIM, first=True, line=1.4)

    if sl.get("callout"):
        callout(slide, y + ch + Inches(0.24), sl["callout"])


def s_pipeline(slide, sl, y):
    steps = sl.get("steps", [])
    n = len(steps) or 1
    rect(slide, M, y, CONTENT_W, Pt(2), fill=BORDER)
    y += Inches(0.3)

    gap = Inches(0.2)
    cw = (CONTENT_W - gap * (n - 1)) / n

    for i, st in enumerate(steps):
        x = M + (cw + gap) * i
        d = Inches(0.42)
        from pptx.enum.shapes import MSO_SHAPE

        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
        circ.fill.solid()
        circ.fill.fore_color.rgb = ACCENT_DEEP
        circ.line.fill.background()
        circ.shadow.inherit = False
        ctf = circ.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        para(ctf, st["n"], 13, WHITE, bold=True, font=MONO, first=True,
             align=PP_ALIGN.CENTER)

        tf = textbox(slide, x, y + Inches(0.58), cw - Inches(0.08), Inches(0.4))
        para(tf, st["name"], 14, TEXT, bold=True, first=True)

        tf = textbox(slide, x, y + Inches(1.02), cw - Inches(0.08), Inches(2.1))
        para(tf, st["body"], 11, DIM, first=True, line=1.4)


def s_compare(slide, sl, y):
    gap = Inches(0.26)
    cw = (CONTENT_W - gap) / 2
    ch = min(Inches(3.5), H - y - Inches(0.9))

    for side, key, colour in ((0, "left", CYAN), (1, "right", ACCENT)):
        d = sl[key]
        x = M + (cw + gap) * side
        rect(slide, x, y, cw, ch, fill=PANEL, line=BORDER)
        rect(slide, x, y, cw, Inches(0.52), fill=PANEL2)
        rect(slide, x, y, Pt(4), Inches(0.52), fill=colour)

        tf = textbox(slide, x + Inches(0.24), y + Inches(0.13), cw - Inches(0.48),
                     Inches(0.36))
        para(tf, d["heading"], 15, colour, bold=True, first=True)

        tf = textbox(slide, x + Inches(0.24), y + Inches(0.68), cw - Inches(0.48),
                     ch - Inches(0.9))
        for i, row in enumerate(d["rows"]):
            para(tf, row, 12, DIM, first=(i == 0), space_after=10, line=1.35)


def s_table(slide, sl, y):
    head, rows = sl["head"], sl["rows"]
    ncol, nrow = len(head), len(rows) + 1
    th = min(Inches(3.5), H - y - Inches(0.9))

    shape = slide.shapes.add_table(nrow, ncol, M, y, CONTENT_W, th)
    tbl = shape.table
    tbl.first_row = True

    widths = ([Inches(1.7), Inches(2.9), CONTENT_W - Inches(4.6)] if ncol == 3
              else [CONTENT_W / ncol] * ncol)
    for i, wd in enumerate(widths[:ncol]):
        tbl.columns[i].width = Emu(int(wd))

    for c, text in enumerate(head):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG
        cell.margin_left = cell.margin_right = Inches(0.1)
        tf = cell.text_frame
        tf.word_wrap = True
        para(tf, text.upper(), 10, ACCENT, bold=True, first=True)

    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r % 2 else BG
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            tf = cell.text_frame
            tf.word_wrap = True
            is_first = c == 0
            para(tf, text, 10.5 if not is_first else 11,
                 TEXT if is_first else DIM,
                 bold=is_first,
                 font=MONO if text.startswith("vllm-sr/") else SANS,
                 first=True, line=1.25)


# -------------------------------------------------------------------- main


def main():
    deck = load_deck()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    total = len(deck["slides"])
    for idx, sl in enumerate(deck["slides"], start=1):
        slide = blank(prs)
        t = sl["type"]

        if t == "title":
            s_title(slide, sl)
        elif t == "closer":
            s_closer(slide, sl)
        else:
            y = heading(slide, sl["title"], sl.get("intro"))
            if t == "statement":
                s_statement(slide, sl, y)
            elif t == "columns":
                s_columns(slide, sl, y)
            elif t == "pipeline":
                s_pipeline(slide, sl, y)
            elif t == "compare":
                s_compare(slide, sl, y)
            elif t == "table":
                s_table(slide, sl, y)
            else:
                sys.exit(f"Unknown slide type {t!r} on slide {idx}")

        chrome(slide, sl.get("source"), idx, total, sl.get("at"))
        notes(slide, sl.get("note"))

    prs.save(OUT)
    print(f"wrote {OUT.relative_to(HERE.parent.parent)}  ({total} slides)")


if __name__ == "__main__":
    main()
